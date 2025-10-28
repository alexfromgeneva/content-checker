import io
import os
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from flask import Flask, jsonify, request
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover - dependency issue should surface during runtime
    raise RuntimeError("python-docx is required for DOCX extraction") from exc


app = Flask(__name__)

SUPPORTED_LANGUAGES: Dict[str, str] = {
    "en": "English",
    "fr": "French",
    "es": "Spanish",
    "ar": "Arabic",
    "zh": "Simplified Chinese",
    "ru": "Russian",
}

ALLOWED_EXTENSIONS = {".txt", ".docx", ".xlsx"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB

# Simplified terminology sample. In production this should be hydrated from UNTERM/WMO sources.
TERMINOLOGY_DB: Tuple[Dict[str, object], ...] = (
    {
        "reference": "https://unterm.un.org/unterm2/en/view/WMO_Atmosphere",
        "category": "meteorology",
        "translations": {
            "en": "atmosphere",
            "fr": "atmosphère",
            "es": "atmósfera",
            "ar": "الغلاف الجوي",
            "zh": "大气层",
            "ru": "атмосфера",
        },
    },
    {
        "reference": "https://unterm.un.org/unterm2/en/view/WMO_ClimateChange",
        "category": "climate",
        "translations": {
            "en": "climate change",
            "fr": "changement climatique",
            "es": "cambio climático",
            "ar": "تغير المناخ",
            "zh": "气候变化",
            "ru": "изменение климата",
        },
    },
    {
        "reference": "https://unterm.un.org/unterm2/en/view/WMO_Precipitation",
        "category": "meteorology",
        "translations": {
            "en": "precipitation",
            "fr": "précipitations",
            "es": "precipitación",
            "ar": "الهطول",
            "zh": "降水",
            "ru": "осадки",
        },
    },
    {
        "reference": "https://unterm.un.org/unterm2/en/view/WMO_EarlyWarning",
        "category": "disaster-risk",
        "translations": {
            "en": "early warning system",
            "fr": "système d'alerte précoce",
            "es": "sistema de alerta temprana",
            "ar": "نظام الإنذار المبكر",
            "zh": "预警系统",
            "ru": "система раннего предупреждения",
        },
    },
    {
        "reference": "https://unterm.un.org/unterm2/en/view/WMO_SustainableDevelopment",
        "category": "sustainability",
        "translations": {
            "en": "sustainable development",
            "fr": "développement durable",
            "es": "desarrollo sostenible",
            "ar": "التنمية المستدامة",
            "zh": "可持续发展",
            "ru": "устойчивое развитие",
        },
    },
)

COUNTRY_TERMS: Tuple[Dict[str, object], ...] = (
    {
        "reference": "https://unterm.un.org/unterm2/en/country/DEU",
        "translations": {
            "en": "Germany",
            "fr": "Allemagne",
            "es": "Alemania",
            "ar": "ألمانيا",
            "zh": "德国",
            "ru": "Германия",
        },
    },
    {
        "reference": "https://unterm.un.org/unterm2/en/country/CAN",
        "translations": {
            "en": "Canada",
            "fr": "Canada",
            "es": "Canadá",
            "ar": "كندا",
            "zh": "加拿大",
            "ru": "Канада",
        },
    },
    {
        "reference": "https://unterm.un.org/unterm2/en/country/ARG",
        "translations": {
            "en": "Argentina",
            "fr": "Argentine",
            "es": "Argentina",
            "ar": "الأرجنتين",
            "zh": "阿根廷",
            "ru": "Аргентина",
        },
    },
)

CATEGORY_WEIGHTS: Dict[str, float] = {
    "terminology": 0.40,
    "grammar": 0.25,
    "style": 0.15,
    "completeness": 0.10,
    "formatting": 0.10,
}

CATEGORY_MAPPING: Dict[str, str] = {
    "terminology": "terminology",
    "consistency": "terminology",
    "country": "terminology",
    "technical": "grammar",
    "grammar": "grammar",
    "style": "style",
    "completeness": "completeness",
    "formatting": "formatting",
}

SEVERITY_DEDUCTIONS: Dict[str, int] = {
    "high": 35,
    "medium": 20,
    "low": 8,
}


@dataclass
class Issue:
    type: str
    severity: str
    issue: str
    suggestion: str
    reference: str = ""
    original: str = ""
    translation: str = ""

    def as_dict(self) -> Dict[str, str]:
        data = asdict(self)
        return {k: v for k, v in data.items() if v}


@dataclass
class TranslationResult:
    overallScore: float
    issues: List[Issue] = field(default_factory=list)
    summary: str = ""
    strengths: List[str] = field(default_factory=list)
    verifiedTerms: List[str] = field(default_factory=list)

    def as_dict(self) -> Dict[str, object]:
        return {
            "overallScore": self.overallScore,
            "issues": [issue.as_dict() for issue in self.issues],
            "summary": self.summary,
            "strengths": self.strengths,
            "verifiedTerms": self.verifiedTerms,
        }


def sanitize_input(text: str) -> str:
    if not text:
        return ""
    clean = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.IGNORECASE | re.DOTALL)
    clean = re.sub(r"<[^>]+>", "", clean)
    return clean.strip()


def detect_numbers(text: str) -> Iterable[str]:
    return re.findall(r"\d+[\d,.]*", text)


def length_ratio(source: str, target: str) -> float:
    if not source or not target:
        return 0.0
    return len(target) / max(len(source), 1)


def extract_text_from_file(file_storage) -> str:
    filename = secure_filename(file_storage.filename or "uploaded")
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension}")

    file_storage.stream.seek(0, os.SEEK_END)
    file_size = file_storage.stream.tell()
    file_storage.stream.seek(0)

    if file_size > MAX_FILE_SIZE:
        raise ValueError("File exceeds maximum size of 10MB")

    data = file_storage.read()
    file_storage.stream.seek(0)

    if extension == ".txt":
        return data.decode("utf-8-sig")

    if extension == ".docx":
        document = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        rows: List[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                if row:
                    rows.append("\t".join(str(cell) for cell in row if cell is not None))
        return "\n".join(rows)

    raise ValueError("Unsupported file type")


def find_terminology_matches(source: str, target: str, source_lang: str, target_lang: str) -> Tuple[List[Issue], List[str]]:
    issues: List[Issue] = []
    verified_terms: List[str] = []
    normalized_source = source.lower()
    normalized_target = target.lower()

    for entry in TERMINOLOGY_DB:
        source_term = entry["translations"].get(source_lang)
        target_term = entry["translations"].get(target_lang)
        if not source_term or not target_term:
            continue

        if source_term.lower() in normalized_source:
            if target_term.lower() in normalized_target:
                verified_terms.append(target_term)
            else:
                issues.append(
                    Issue(
                        type="terminology",
                        severity="high",
                        issue=f"Missing official translation for '{source_term}'.",
                        suggestion=f"Use '{target_term}' as recommended by UNTERM/WMO.",
                        reference=entry["reference"],
                        original=source_term,
                    )
                )

    for entry in COUNTRY_TERMS:
        source_term = entry["translations"].get(source_lang)
        target_term = entry["translations"].get(target_lang)
        if not source_term or not target_term:
            continue

        if source_term.lower() in normalized_source and target_term.lower() not in normalized_target:
            issues.append(
                Issue(
                    type="country",
                    severity="high",
                    issue=f"Country name '{source_term}' must be translated using the official UN form.",
                    suggestion=f"Use '{target_term}'.",
                    reference=entry["reference"],
                    original=source_term,
                )
            )

    return issues, verified_terms


def check_numeric_consistency(source: str, target: str) -> List[Issue]:
    issues: List[Issue] = []
    source_numbers = set(detect_numbers(source))
    target_numbers = set(detect_numbers(target))
    missing = source_numbers - target_numbers
    if missing:
        for number in sorted(missing):
            issues.append(
                Issue(
                    type="technical",
                    severity="medium",
                    issue=f"Numeric value '{number}' is missing in the translation.",
                    suggestion="Ensure all numeric data is preserved.",
                    original=number,
                )
            )
    return issues


def check_formatting(target: str) -> List[Issue]:
    issues: List[Issue] = []

    if "  " in target:
        issues.append(
            Issue(
                type="style",
                severity="low",
                issue="Extra spacing detected.",
                suggestion="Reduce double spaces to a single space.",
            )
        )

    parentheses_balance = target.count("(") == target.count(")")
    if not parentheses_balance:
        issues.append(
            Issue(
                type="formatting",
                severity="medium",
                issue="Unbalanced parentheses detected.",
                suggestion="Review punctuation and ensure parentheses are balanced.",
            )
        )

    if re.search(r"\S{25,}", target):
        issues.append(
            Issue(
                type="style",
                severity="low",
                issue="Long unbroken string detected (possible formatting issue).",
                suggestion="Insert appropriate spaces or hyphenation.",
            )
        )

    if any(char.isdigit() for char in target) and not re.search(r"\d", target[-5:]):
        pass  # digits appear somewhere; handled by numeric consistency

    return issues


def evaluate_completeness(source: str, target: str) -> List[Issue]:
    issues: List[Issue] = []
    ratio = length_ratio(source, target)
    if not target.strip():
        issues.append(
            Issue(
                type="completeness",
                severity="high",
                issue="Translation text is empty.",
                suggestion="Provide a translation before running checks.",
            )
        )
        return issues

    if ratio < 0.5 or ratio > 2.0:
        issues.append(
            Issue(
                type="completeness",
                severity="medium",
                issue=f"Length ratio between source and translation is {ratio:.2f}, outside expected range (0.5 - 2.0).",
                suggestion="Review for missing or added content.",
            )
        )
    return issues


def compute_score(issues: List[Issue]) -> float:
    subscores: Dict[str, float] = {category: 100.0 for category in CATEGORY_WEIGHTS}

    for issue in issues:
        category = CATEGORY_MAPPING.get(issue.type)
        if not category:
            continue
        deduction = SEVERITY_DEDUCTIONS.get(issue.severity, 10)
        subscores[category] = max(0.0, subscores[category] - deduction)

    weighted_sum = sum(subscores[category] * weight for category, weight in CATEGORY_WEIGHTS.items())
    total_weight = sum(CATEGORY_WEIGHTS.values())
    return round(weighted_sum / total_weight, 1)


def summarize_results(result: TranslationResult, source: str, target: str) -> None:
    high_severity = [issue for issue in result.issues if issue.severity == "high"]
    medium_severity = [issue for issue in result.issues if issue.severity == "medium"]

    if high_severity:
        summary = f"High-severity issues found: {len(high_severity)}. Address these before publication."
    elif medium_severity:
        summary = "Translation is generally acceptable but review medium-severity issues for improvement."
    else:
        summary = "Translation aligns with UN/WMO expectations based on the available checks."

    if result.verifiedTerms:
        result.strengths.append(
            f"Verified official terminology: {', '.join(sorted(set(result.verifiedTerms)))}."
        )

    if not detect_numbers(source) or not any(issue.type == "technical" for issue in result.issues):
        result.strengths.append("All numeric data preserved.")

    if result.overallScore >= 85:
        result.strengths.append("High overall quality score (≥85).")

    result.summary = summary


def check_translation(source: str, target: str, source_lang: str, target_lang: str) -> TranslationResult:
    source_clean = sanitize_input(source)
    target_clean = sanitize_input(target)

    issues: List[Issue] = []
    terminology_issues, verified_terms = find_terminology_matches(
        source_clean, target_clean, source_lang, target_lang
    )
    issues.extend(terminology_issues)

    issues.extend(check_numeric_consistency(source_clean, target_clean))
    issues.extend(check_formatting(target_clean))
    issues.extend(evaluate_completeness(source_clean, target_clean))

    score = compute_score(issues)
    result = TranslationResult(overallScore=score, issues=issues, verifiedTerms=verified_terms)
    summarize_results(result, source_clean, target_clean)
    return result


@app.route("/api/health", methods=["GET"])
def health_check():
    return jsonify({"status": "ok", "supportedLanguages": SUPPORTED_LANGUAGES})


@app.route("/api/extract-text", methods=["POST"])
def extract_text_endpoint():
    if "file" not in request.files:
        return jsonify({"error": "File is required"}), 400

    file_storage = request.files["file"]
    try:
        text = extract_text_from_file(file_storage)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    return jsonify({"fileName": file_storage.filename, "text": text})


@app.route("/api/check-translation", methods=["POST"])
def check_translation_endpoint():
    try:
        payload = request.get_json(force=True)
    except (TypeError, ValueError):
        return jsonify({"error": "Invalid JSON payload"}), 400

    source_text = sanitize_input(payload.get("sourceText", ""))
    source_language = payload.get("sourceLanguage", "en").lower()
    translations = payload.get("translations", {})

    if source_language not in SUPPORTED_LANGUAGES:
        return jsonify({"error": f"Unsupported source language '{source_language}'"}), 400

    if not isinstance(translations, dict) or not translations:
        return jsonify({"error": "translations object with at least one target language is required"}), 400

    results: Dict[str, Dict[str, object]] = {}
    for language_code, target_text in translations.items():
        lang_code = language_code.lower()
        if lang_code not in SUPPORTED_LANGUAGES:
            results[lang_code] = {"error": f"Unsupported target language '{language_code}'"}
            continue

        result = check_translation(source_text, target_text, source_language, lang_code)
        results[lang_code] = result.as_dict()

    return jsonify(
        {
            "sourceLanguage": source_language,
            "results": results,
        }
    )


@app.route("/api/check-translation/batch", methods=["POST"])
def batch_check_translation():
    try:
        payload = request.get_json(force=True)
    except (TypeError, ValueError):
        return jsonify({"error": "Invalid JSON payload"}), 400

    items = payload.get("items", [])
    if not isinstance(items, list) or not items:
        return jsonify({"error": "items array is required"}), 400

    batch_results: List[Dict[str, object]] = []
    for item in items:
        file_name = item.get("fileName", "")
        source_text = item.get("sourceText", "")
        source_language = item.get("sourceLanguage", "en").lower()
        translations = item.get("translations", {})

        if source_language not in SUPPORTED_LANGUAGES:
            batch_results.append(
                {
                    "fileName": file_name,
                    "error": f"Unsupported source language '{source_language}'",
                }
            )
            continue

        if not translations:
            batch_results.append(
                {
                    "fileName": file_name,
                    "error": "No translations provided",
                }
            )
            continue

        file_result: Dict[str, object] = {
            "fileName": file_name,
            "sourceText": sanitize_input(source_text),
            "sourceLang": source_language,
            "results": {},
        }

        for language_code, translation_text in translations.items():
            lang_code = language_code.lower()
            if lang_code not in SUPPORTED_LANGUAGES:
                file_result["results"][lang_code] = {"error": f"Unsupported target language '{language_code}'"}
                continue

            result = check_translation(source_text, translation_text, source_language, lang_code)
            file_result["results"][lang_code] = result.as_dict()

        batch_results.append(file_result)

    return jsonify({"batchResults": batch_results})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
